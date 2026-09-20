from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = Path('docs/bioinfo_teaching_system_presentation.pptx')
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = '10243A'; BLUE = '2E75B6'; TEAL = '00A6A6'; ORANGE = 'F39C12'; PURPLE = '7030A0'; RED = 'C0504D'; DARK = '1F2937'; GRAY = '64748B'; LIGHT = 'F6F8FB'; WHITE = 'FFFFFF'
FONT = 'Noto Sans CJK SC'


def rgb(hexcode): return RGBColor.from_string(hexcode)

def set_bg(slide, color=WHITE):
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = rgb(color)

def text_box(slide, left, top, width, height, text='', size=22, color=DARK, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(0.04); tf.margin_right = Inches(0.04); tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text; run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    return shape

def title(slide, text, dark=False):
    color = WHITE if dark else NAVY
    text_box(slide, .72, .42, 11.9, .6, text, 28, color, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.75), Inches(1.13), Inches(1.12), Inches(.055))
    line.fill.solid(); line.fill.fore_color.rgb = rgb(ORANGE); line.line.fill.background()

def footer(slide, number, dark=False):
    color = 'CBD5E1' if dark else '94A3B8'
    text_box(slide, .75, 7.08, 10.5, .18, '基于AI大模型的生物信息学课程教学一体化平台 · 项目汇报', 8.5, color)
    text_box(slide, 12.05, 7.05, .45, .2, str(number), 9, color, True, PP_ALIGN.RIGHT)

def bullet_slide(number, heading, bullets, accent=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide); title(slide, heading)
    y = 1.52
    for bullet in bullets:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(.86), Inches(y+.09), Inches(.13), Inches(.13))
        dot.fill.solid(); dot.fill.fore_color.rgb = rgb(accent); dot.line.fill.background()
        text_box(slide, 1.15, y, 11.3, .54, bullet, 16, DARK)
        y += .78
    footer(slide, number)
    return slide

def flow_slide(number, heading, steps, caption, accent=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); set_bg(slide); title(slide, heading)
    n = len(steps); left = .75; gap = .18; width = (11.85 - gap * (n - 1)) / n
    for i, (label, detail, color) in enumerate(steps):
        x = left + i * (width + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.05), Inches(width), Inches(1.75))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color); shape.line.fill.background()
        tf = shape.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label; r.font.name = FONT; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = rgb(WHITE)
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = detail; r2.font.name = FONT; r2.font.size = Pt(12); r2.font.color.rgb = rgb(WHITE)
        if i < n - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + width + .02), Inches(2.72), Inches(.14), Inches(.28))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = rgb('94A3B8'); arrow.line.fill.background()
    text_box(slide, 1.0, 4.65, 11.3, .7, caption, 17, GRAY, False, PP_ALIGN.CENTER)
    footer(slide, number)
    return slide

def pill(slide, left, top, width, text, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(.55))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color); shape.line.fill.background()
    tf=shape.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=text;r.font.name=FONT;r.font.size=Pt(16);r.font.bold=True;r.font.color.rgb=rgb(WHITE)

# 1 Cover
slide=prs.slides.add_slide(prs.slide_layouts[6]);set_bg(slide,NAVY)
text_box(slide,.9,1.45,11.5,1.0,'基于AI大模型的生物信息学课程教学一体化平台',32,WHITE,True,PP_ALIGN.CENTER)
text_box(slide,1.6,2.72,10.1,.7,'课程教学 · 智能出题 · 知识库 · 考试评价',20,'D9E4F0',False,PP_ALIGN.CENTER)
for x,t,c in [(2.05,'课程资源',BLUE),(4.35,'AI 出题',TEAL),(6.65,'人工审核',ORANGE),(8.95,'学习评价',PURPLE)]: pill(slide,x,4.15,1.85,t,c)
text_box(slide,1.5,6.25,10.4,.28,'项目汇报展示 · 2026年08月26日 · 汇报成员：[待确认]',12,'CBD5E1',False,PP_ALIGN.CENTER)
footer(slide,1,True)

# 2
bullet_slide(2,'目录',['项目背景与系统定位','系统架构与技术栈','核心模块：学生、教师、管理员','知识库、知识图谱与文献检索','功能演示、技术亮点与后续工作'],TEAL)
# 3
bullet_slide(3,'项目概述：面向课程教学的完整闭环',['面向生物信息学课程的 Web 教学、练习与考试平台','学生完成练习、考试、成绩查看与学习建议','教师使用 AI 生成题目，并通过审核保障题目质量','管理员管理账户、题目与知识库资源','资源 → 出题 → 审核 → 学习/考试 → 评价反馈'],BLUE)
# 4 Architecture
slide=prs.slides.add_slide(prs.slide_layouts[6]);set_bg(slide);title(slide,'系统架构：Flask + MySQL + AI 服务')
for x,t,c in [(0.8,'浏览器\nJinja2 模板',BLUE),(3.5,'Flask\napp.py',TEAL),(6.2,'MySQL\n业务数据',ORANGE),(8.9,'AI 服务\nDeepSeek / Ollama',PURPLE),(11.0,'文献 API\nEurope PMC / Crossref',RED)]:
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.5),Inches(1.7),Inches(1.25));shape.fill.solid();shape.fill.fore_color.rgb=rgb(c);shape.line.fill.background();tf=shape.text_frame;tf.clear();tf.vertical_anchor=MSO_ANCHOR.MIDDLE;p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=t;r.font.name=FONT;r.font.size=Pt(16);r.font.bold=True;r.font.color.rgb=rgb(WHITE)
for x in [2.62,5.32,8.02,10.15]:
    a=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(2.92),Inches(.58),Inches(.35));a.fill.solid();a.fill.fore_color.rgb=rgb(GRAY);a.line.fill.background()
text_box(slide,1.0,4.5,11.3,1.15,'表现层：Jinja2、HTML/CSS、原生 JavaScript、D3.js\n应用层：登录、题目、任务、知识库、组卷、考试、API 路由集中于 Flask 应用\n数据层：用户、题目、任务、文档、概念、试卷和学习记录',18,DARK)
footer(slide,4)
# 5
bullet_slide(5,'技术栈总览',['后端：Flask、Flask-Session、Gunicorn','数据库：MySQL + mysql-connector-python','AI 与网络：DeepSeek、Ollama、requests、aiohttp','文档处理：PyPDF2、python-docx、jieba','前端：Jinja2、原生 JavaScript、D3.js'],TEAL)
#6
bullet_slide(6,'学生模块：练习、考试与反馈',['按主题、语言与难度抽取编程题或理论题','Python、Perl、Java、C 代码编辑、运行与提交','自动评分写入练习/考试记录，返回分数和评语','成绩页汇总 AI 评分、教师评分与考试记录','学习建议聚合考试、练习及知识库学习行为'],BLUE)
flow_slide(7,'教师出题：从需求到待审核', [('输入','主题 · 难度 · 语言',BLUE),('生成','DeepSeek / Ollama',TEAL),('任务','进度 · 日志 · 状态',PURPLE),('落库','questions / theory_questions',ORANGE),('待审','教师质量把关',RED)], '理论题支持 PDF 或知识库文本；单次文本最多 50,000 字符。', PURPLE)
flow_slide(8,'题目审核：质量控制链', [('待审核','统一题目列表',ORANGE),('查看','题目 · 答案 · 关联文档',BLUE),('决策','通过 / 驳回',TEAL),('记录','原因 · 审核人 · 时间',PURPLE),('复用','题库 · 考试 · 组卷',RED)], '三状态：pending → approved / rejected。', ORANGE)
flow_slide(9,'知识库：文档到知识图谱', [('上传','PDF · TXT · DOCX',BLUE),('抽取','概念 · 类型 · 频率',TEAL),('建边','概念关系 · 置信度',PURPLE),('聚合','合并同名概念',ORANGE),('可视化','D3 力导向图',RED)], '知识库把文档内容转化为可浏览、可关联的学习资源。', TEAL)
flow_slide(10,'智能组卷：题库到 A/B 卷', [('筛选','已审核题目',BLUE),('配置','题型 · 数量 · 分值',TEAL),('平衡','难度 · 知识点覆盖',PURPLE),('生成','A 卷 · B 卷 · 答案',ORANGE),('输出','预览 · PDF · 克隆',RED)], '题目不足时给出降级提示，不静默生成空试卷。', RED)
flow_slide(11,'文献检索：检索到学习沉淀', [('检索','期刊 · 关键词 · 年份',BLUE),('获取','Europe PMC',TEAL),('回退','无记录 → Crossref',PURPLE),('分析','创新点 · 深入研读',ORANGE),('沉淀','CSV · 学习笔记',RED)], '学生与教师均可访问文献检索页面；教师/管理员可维护期刊。', PURPLE)
#12 demo
slide=prs.slides.add_slide(prs.slide_layouts[6]);set_bg(slide);title(slide,'功能演示主线')
steps=[('1','知识库/PDF\n生成理论题',BLUE),('2','教师审核\n进入题库',TEAL),('3','学生练习\n运行并提交',ORANGE),('4','AI 评分\n学习反馈',PURPLE),('5','组卷与\n教学复用',RED)]
for i,(n,t,c) in enumerate(steps):
    x=.85+i*2.48
    circle=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+.57),Inches(2.2),Inches(.58),Inches(.58));circle.fill.solid();circle.fill.fore_color.rgb=rgb(c);circle.line.fill.background();tf=circle.text_frame;tf.clear();tf.vertical_anchor=MSO_ANCHOR.MIDDLE;p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=n;r.font.name=FONT;r.font.size=Pt(18);r.font.bold=True;r.font.color.rgb=rgb(WHITE)
    text_box(slide,x,3.05,1.75,.75,t,16,DARK,True,PP_ALIGN.CENTER)
    if i<4:
        a=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+1.85),Inches(2.35),Inches(.55),Inches(.25));a.fill.solid();a.fill.fore_color.rgb=rgb('CBD5E1');a.line.fill.background()
text_box(slide,1.05,4.65,11.2,.95,'建议现场演示：教师生成题目 → 审核通过 → 学生完成练习 → 查看知识库/图谱/笔记 → 教师组卷。',18,GRAY,False,PP_ALIGN.CENTER)
footer(slide,12)
#13
bullet_slide(13,'核心数据模型与数据流',['users：账号、工号、角色、登录状态','questions / theory_questions：编程题与理论题','tasks：AI 生成任务、进度、日志与结果','knowledge_bases → documents → concepts → relations','exams / exercises / suggestions：学习记录、评价与反馈'],BLUE)
#14
bullet_slide(14,'技术实现亮点',['DeepSeek 与本地 Ollama 可切换，用于出题、评分与建议','多语言代码执行覆盖 Python、Perl、Java 与 C','同名概念合并、频率聚合、关系端点映射支撑图谱展示','旧数据库字段采用幂等补齐，降低历史数据升级成本','文献检索具备 Europe PMC → Crossref 的容错回退'],TEAL)
#15
bullet_slide(15,'当前状态与后续工作',['已实现三角色控制台、AI 出题/审核、知识库/图谱、练习/考试、组卷、文献检索和学习笔记','当前主应用约 9,677 行 app.py，模板约 54 个，采用单体 Flask 架构','后续统一项目文档、部署配置与数据库 schema 版本说明','后续强化代码执行隔离、权限控制、自动化测试与任务队列','团队成员分工：[待确认]'],RED)
#16
slide=prs.slides.add_slide(prs.slide_layouts[6]);set_bg(slide,NAVY)
text_box(slide,.9,1.75,11.5,.72,'感谢聆听',34,WHITE,True,PP_ALIGN.CENTER)
text_box(slide,1.5,3.0,10.3,.45,'Q&A',24,'D9E4F0',True,PP_ALIGN.CENTER)
text_box(slide,1.2,4.1,10.9,.75,'当前部署入口：http://172.24.64.53:8080\n项目链接：[待补充]',16,'CBD5E1',False,PP_ALIGN.CENTER)
footer(slide,16,True)

prs.save(OUT)
print(OUT)
